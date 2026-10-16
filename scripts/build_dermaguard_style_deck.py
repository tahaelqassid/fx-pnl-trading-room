"""
DermaGuard-Style Human Presentation Generator for PFA Defense.
Matches the exact visual layout, card structure, and style of Mohamed Taha's previous UIR presentation:
- Clean Light Canvas (#FFFFFF / #F8FAFC)
- Breadcrumb slide headers: '01 / The Problem', '02 / System Architecture', etc.
- Metric callouts, colored badges, structured flow with arrows
- 100% English, ZERO gold references, fully focused on the FX P&L project
- Dedicated logo placeholders for UIR and DXC
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Root directories
ROOT_DIR = Path(__file__).resolve().parent.parent
PRES_DIR = ROOT_DIR / "presentation"
PRES_DIR.mkdir(parents=True, exist_ok=True)

# DermaGuard Color Palette (Clean Minimalist Light Mode)
BG_PAGE = RGBColor(248, 250, 252)       # #F8FAFC Light slate base
CARD_BG = RGBColor(255, 255, 255)       # Pure white cards
BORDER_DEFAULT = RGBColor(226, 232, 240)# #E2E8F0 Subtle border
BORDER_ACTIVE = RGBColor(203, 213, 225) # #CBD5E1

# Accent Colors
C_NAVY = RGBColor(15, 23, 42)           # #0F172A Primary text
C_MUTED = RGBColor(100, 116, 139)       # #64748B Secondary text
C_BLUE = RGBColor(37, 99, 235)          # #2563EB Royal Blue
C_EMERALD = RGBColor(16, 185, 129)      # #10B981 Green
C_AMBER = RGBColor(217, 119, 6)         # #D97706 Gold / Amber
C_ROSE = RGBColor(225, 29, 72)          # #E11D48 Red / Rose
C_PURPLE = RGBColor(126, 34, 206)       # #7E22CE Purple
C_TEAL = RGBColor(13, 148, 136)         # #0D9488 Teal

FONT_MAIN = "Arial"
FONT_MONO = "Consolas"


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_PAGE

    def add_card(slide, left, top, width, height, bg=CARD_BG, border=BORDER_DEFAULT):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        card.line.color.rgb = border
        card.line.width = Pt(1.1)
        return card

    def add_logo_box(slide, left, top, width, height, text_str="[ Upload Logo Here ]"):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        box.line.color.rgb = RGBColor(203, 213, 225)
        box.line.width = Pt(1.2)
        box.line.dash_style = 2  # Dashed line
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text_str
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        return box

    def add_header(slide, number_str, title_str):
        # DermaGuard Breadcrumb style: '01  /  The Problem'
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = f"{number_str}  /  {title_str}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = C_NAVY

        # Subtle logo placeholders at top-right
        add_logo_box(slide, Inches(10.5), Inches(0.4), Inches(1.0), Inches(0.6), "[ UIR ]")
        add_logo_box(slide, Inches(11.6), Inches(0.4), Inches(1.0), Inches(0.6), "[ DXC ]")

    def set_notes(slide, notes_txt):
        ns = slide.notes_slide
        tf = ns.notes_text_frame
        tf.text = notes_txt.strip()

    # =========================================================================
    # SLIDE 1: COVER SLIDE (DERMAGUARD TITLE STYLE)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Top Tag
    tb_tag = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.4))
    tf_tag = tb_tag.text_frame
    tf_tag.word_wrap = True
    p = tf_tag.paragraphs[0]
    p.text = "UIR · School of Data Science & Digital Engineering · Data Engineering (4th Year) · 2025–2026"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    # Title & Subtitle Card
    title_card = add_card(s1, Inches(0.8), Inches(1.1), Inches(8.5), Inches(2.2))
    tb_t = s1.shapes.add_textbox(Inches(1.1), Inches(1.25), Inches(8.0), Inches(1.9))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    
    p = tf_t.paragraphs[0]
    p.text = "FX Trading Room P&L Platform"
    p.font.name = FONT_MAIN
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p_sub = tf_t.add_paragraph()
    p_sub.text = "Data Engineering Platform for Foreign Exchange Profit & Loss Monitoring, Bank Al-Maghrib Compliance & Front-to-Back Reconciliation (Accounting Currency: MAD)"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = C_MUTED
    p_sub.space_before = Pt(8)

    # Right: Big Logo Placement Box
    add_logo_box(s1, Inches(9.6), Inches(0.6), Inches(2.9), Inches(2.7), "[ Drag & Drop UIR / DXC Logos Here ]")

    # 4 Metadata Cards in a Row (Exact DermaGuard Style)
    meta_items = [
        ("🏦", "Accounting Currency", "MAD (Moroccan Dirham)\nBank Al-Maghrib Fixing", C_BLUE),
        ("⚙️", "Data Pipeline", "PostgreSQL · dbt Core\nApache Airflow (15:35 EOD)", C_EMERALD),
        ("📊", "Trading Engine", "Python 3.11 · FIFO & CMP\nMulti-Desk Portfolio Keeping", C_AMBER),
        ("🛡️", "Reconciliation (OS6)", "Front-to-Back Matching\nSWIFT MT300 Dispute Hub", C_PURPLE)
    ]
    card_w = Inches(2.75)
    for i, (icon, h_txt, b_txt, col) in enumerate(meta_items):
        cx = Inches(0.8 + i * 2.95)
        add_card(s1, cx, Inches(3.55), card_w, Inches(2.2))
        tb_m = s1.shapes.add_textbox(cx + Inches(0.2), Inches(3.7), card_w - Inches(0.4), Inches(1.9))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        
        p = tf_m.paragraphs[0]
        p.text = f"{icon}  {h_txt}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

        p_b = tf_m.add_paragraph()
        p_b.text = b_txt
        p_b.font.name = FONT_MAIN
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = C_NAVY
        p_b.space_before = Pt(8)

    # Bottom Authors Bar
    bot_card = add_card(s1, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), bg=RGBColor(241, 245, 249))
    tb_b = s1.shapes.add_textbox(Inches(1.1), Inches(6.1), Inches(11.1), Inches(0.7))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p1 = tf_b.paragraphs[0]
    p1.text = "Author: Mohamed Taha El Qassid (Data Engineering, 4th Year)   |   Host Company: DXC Technology Morocco"
    p1.font.name = FONT_MAIN
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY
    p2 = tf_b.add_paragraph()
    p2.text = "Supervisors: Prof. Fayçal Zinedine (UIR) · Samir Gourroum (DXC Technology) · PFA S8 Final Defense"
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(10)
    p2.font.color.rgb = C_MUTED
    p2.space_before = Pt(2)

    set_notes(s1, "Good morning esteemed jury members, Prof. Zinedine, and Mr. Gourroum. I am Mohamed Taha El Qassid, 4th-year Data Engineering student at UIR. Today I present my final year project carried out at DXC Technology: a data engineering platform for FX P&L calculation, Bank Al-Maghrib compliance, and Front-to-Back reconciliation.")

    # =========================================================================
    # SLIDE 2: 01 / THE PROBLEM (EXACT DERMAGUARD METRICS GRID)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "01", "The Problem")

    # 4 Big Metric Cards (DermaGuard style)
    p_cards = [
        ("100s of Millions", "MAD Daily Traded Volume", "Corporate flows from OCP Group (phosphate export), Renault Tanger, and Afriquia (crude oil imports).", C_BLUE),
        ("15:30 EOD", "BAM Mandatory Fixing", "Moroccan banks cannot carry unvalued foreign currency books overnight. Daily revaluation to MAD is mandated by Bank Al-Maghrib.", C_AMBER),
        ("4 Fatal Breaks", "Manual Excel Vulnerabilities", "Rate slippage, nominal mismatch, unconfirmed SWIFT affirmations, and rogue trades cause hidden financial leakage.", C_ROSE),
        ("100% Identity", "Accounting Rule BR3", "Total MTM P&L must strictly equal Realized + Unrealized P&L. Zero financial leakage allowed in production.", C_EMERALD)
    ]
    for i, (big_num, lbl, desc, col) in enumerate(p_cards):
        col_idx = i % 2
        row_idx = i // 2
        x = Inches(0.8 + col_idx * 5.95)
        y = Inches(1.5 + row_idx * 2.2)
        
        add_card(s2, x, y, Inches(5.75), Inches(2.0))
        tb = s2.shapes.add_textbox(x + Inches(0.3), y + Inches(0.18), Inches(5.15), Inches(1.65))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = big_num
        p.font.name = FONT_MAIN
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = col

        p_lbl = tf.add_paragraph()
        p_lbl.text = lbl
        p_lbl.font.name = FONT_MAIN
        p_lbl.font.size = Pt(12)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = C_NAVY
        p_lbl.space_before = Pt(2)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_MAIN
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = C_MUTED
        p_desc.space_before = Pt(4)

    # Bottom Punchline Card
    b_card = add_card(s2, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), bg=RGBColor(241, 245, 249))
    tb_b = s2.shapes.add_textbox(Inches(1.1), Inches(6.15), Inches(11.1), Inches(0.6))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "FX Trading Room Platform: Automated EOD revaluation against Bank Al-Maghrib fixing, automated Front-to-Back dispute detection, and guaranteed zero financial leakage."
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    set_notes(s2, "In a trading room, millions of Dirhams are exchanged daily for major Moroccan corporations. Bank Al-Maghrib mandates daily revaluation into MAD at 15:30. Legacy Excel models suffer from rate slippage, nominal breaks, and lack of audit trails. Our platform automates the entire flow.")

    # =========================================================================
    # SLIDE 3: 02 / SYSTEM ARCHITECTURE (DERMAGUARD PIPELINE FLOW)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "02", "System Architecture")

    # Left: The Step-by-Step Flow with Arrows
    left_card = add_card(s3, Inches(0.8), Inches(1.5), Inches(8.3), Inches(5.4))
    tb_f = s3.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(7.7), Inches(5.0))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True

    steps = [
        ("📥  1. Ingestion Layer", "Raw CSV trade tickets, interbank feeds & Bank Al-Maghrib official fixing rates."),
        ("🥉  2. Bronze Staging", "PostgreSQL relational tables (stg_trade, stg_market_rate) with transparent SQLite fallback."),
        ("🥈  3. Silver Marts (dbt)", "dbt Core SQL transformation models (fact_trade_mtm) and synthetic EUR/USD cross conversions."),
        ("🥇  4. Gold Calculation Engine", "Python FIFO lot unwinding deque, WAC rolling average cost, and Bank Al-Maghrib peg checks."),
        ("🔍  5. Front-to-Back Reconciliation", "Automated matching engine detecting rate slippage, nominal mismatches, and unconfirmed SWIFT deals."),
        ("✅  6. Persistence & UI", "Audited daily_pnl table (enforced by CHECK rule), Streamlit trading terminal & 1-click Excel export.")
    ]
    for idx, (title_s, desc_s) in enumerate(steps):
        p = tf_f.paragraphs[0] if idx == 0 else tf_f.add_paragraph()
        p.text = title_s
        p.font.name = FONT_MAIN
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = C_BLUE if "1" in title_s or "2" in title_s else (C_AMBER if "3" in title_s else C_EMERALD)
        if idx > 0:
            p.space_before = Pt(8)
            
        p_d = tf_f.add_paragraph()
        p_d.text = desc_s
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = C_NAVY
        p_d.space_before = Pt(2)

    # Right: Tech Stack Card (DermaGuard Style)
    right_card = add_card(s3, Inches(9.3), Inches(1.5), Inches(3.233), Inches(5.4))
    tb_st = s3.shapes.add_textbox(Inches(9.5), Inches(1.7), Inches(2.833), Inches(5.0))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True

    p = tf_st.paragraphs[0]
    p.text = "Tech Stack"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    stack_pills = [
        "PostgreSQL 15", "dbt Core 1.8", "Apache Airflow", "Python 3.11",
        "Streamlit", "Plotly", "openpyxl", "pytest", "Docker"
    ]
    for item in stack_pills:
        p = tf_st.add_paragraph()
        p.text = f"▸ {item}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        p.space_before = Pt(8)

    set_notes(s3, "Here is our Medallion data engineering architecture: Bronze for raw ingestion into PostgreSQL, Silver for dbt SQL transformations including cross-rate conversion, and Gold for the Python calculation engine. The right panel shows our industrial tech stack.")

    # =========================================================================
    # SLIDE 4: 03 / FINANCIAL CORE — THE 4 BENCHMARK TRADES (+79,460 MAD)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "03", "Financial Core — The 4 Benchmark Trades")

    # 3 Big KPI Cards at Top
    kpis = [
        ("+79,460 MAD", "Total Mark-to-Market P&L", "Overall portfolio profit at BAM closing fixing", C_BLUE),
        ("+20,000 MAD", "Realized P&L (Cash)", "Actual closed gain (unwinding T1 with T2)", C_EMERALD),
        ("+59,460 MAD", "Unrealized P&L (Latent)", "Paper gain on open overnight contracts", C_AMBER)
    ]
    kw = Inches(3.7)
    for i, (num_s, h_s, sub_s, col) in enumerate(kpis):
        kx = Inches(0.8 + i * 4.0)
        add_card(s4, kx, Inches(1.5), kw, Inches(1.5))
        tb = s4.shapes.add_textbox(kx + Inches(0.25), Inches(1.65), kw - Inches(0.5), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = num_s
        p.font.name = FONT_MAIN
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = col
        p_h = tf.add_paragraph()
        p_h.text = h_s
        p_h.font.name = FONT_MAIN
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = C_NAVY
        p_h.space_before = Pt(2)

    # Bottom Table: Exact Math from Report
    bot_card = add_card(s4, Inches(0.8), Inches(3.2), Inches(11.733), Inches(3.7))
    tb_t = s4.shapes.add_textbox(Inches(1.1), Inches(3.35), Inches(11.1), Inches(3.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True

    p = tf_t.paragraphs[0]
    p.text = "Mathematical Breakdown of the 4 Reference Trades (Report Table 2.2 – 2.4)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    t_rows = [
        ("Trade 1 (T1) · USD/MAD Buy 1M @ 9.9200", "BAM Fixing: 9.9500 MAD", "(9.9500 - 9.9200) * 1,000,000 = +30,000 MAD (MTM)"),
        ("Trade 2 (T2) · USD/MAD Sell 500k @ 9.9600", "Unwinds half of T1", "Realized Cash: (9.96 - 9.92) * 500k = +20,000 MAD  |  Remaining Latent: +15,000 MAD"),
        ("Trade 3 (T3) · EUR/MAD Buy 800k @ 10.7700", "BAM Fixing: 10.8000 MAD", "(10.8000 - 10.7700) * 800,000 = +24,000 MAD (Latent)"),
        ("Trade 4 (T4) · EUR/USD Buy 600k @ 1.0820", "Synthetic Cross in MAD", "(10.8000 - (1.0820 * 9.9500)) * 600,000 = +20,460 MAD (Latent)")
    ]
    for name, fix_s, math_s in t_rows:
        p = tf_t.add_paragraph()
        p.text = f"• {name}  [{fix_s}]"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_NAVY
        p.space_before = Pt(8)
        p_m = tf_t.add_paragraph()
        p_m.text = f"  Calculation: {math_s}"
        p_m.font.name = FONT_MONO
        p_m.font.size = Pt(9.5)
        p_m.font.color.rgb = C_BLUE

    set_notes(s4, "Here is our exact mathematical reproduction of the PFE brief benchmark: four trades across USD, EUR, and EUR/USD cross. Total portfolio MTM is exactly +79,460 MAD, decomposed into +20,000 MAD Realized cash and +59,460 MAD Latent paper gain. Rule BR3 holds 100%.")

    # =========================================================================
    # SLIDE 5: 04 / INVENTORY ACCOUNTING — FIFO VS CMP & TYPO AUDIT
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "04", "Inventory Accounting — FIFO vs. CMP & Typo Fix")

    col_w = Inches(5.75)
    col_h = Inches(3.6)

    # Left: FIFO
    c_fifo = add_card(s5, Inches(0.8), Inches(1.5), col_w, col_h)
    tb = s5.shapes.add_textbox(Inches(1.05), Inches(1.7), col_w - Inches(0.5), col_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FIFO Method (Murex / Capital Markets Standard)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = "Position: 2 Purchases (500k @ 9.91 and 500k @ 9.93) + 1 Sale (600k @ 9.95)\n"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "1. Lot 1 exhausted first (500,000 @ 9.91):\n   (9.9500 - 9.9100) * 500,000 = +20,000 MAD\n\n2. Lot 2 partially consumed (100,000 @ 9.93):\n   (9.9500 - 9.9300) * 100,000 = +2,000 MAD\n\nTotal Realized P&L = +22,000 MAD (Remaining: 400k @ 9.93)"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NAVY
    p.space_before = Pt(4)

    # Right: CMP
    c_cmp = add_card(s5, Inches(6.8), Inches(1.5), col_w, col_h)
    tb = s5.shapes.add_textbox(Inches(7.05), Inches(1.7), col_w - Inches(0.5), col_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CMP / WAC (Moroccan Accounting Standard)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Weighted Average Cost Calculation:\n"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "• Weighted Unit Cost:\n   ((500k * 9.91) + (500k * 9.93)) / 1,000,000 = 9.9200 MAD\n\n• Realized P&L on 600,000 Sale:\n   (9.9500 - 9.9200) * 600,000 = +18,000 MAD\n\nTotal Realized P&L = +18,000 MAD (Remaining: 400k @ 9.92)"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NAVY
    p.space_before = Pt(4)

    # Bottom Discovery Card (Exact DermaGuard callout box style)
    disc_card = add_card(s5, Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.6), bg=RGBColor(254, 243, 199), border=RGBColor(245, 158, 11))
    tb_d = s5.shapes.add_textbox(Inches(1.1), Inches(5.45), Inches(11.1), Inches(1.3))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True

    p = tf_d.paragraphs[0]
    p.text = "🔍  KEY AUDIT DISCOVERY & TYPO FIX (Report Section 8.2)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_AMBER

    p_body = tf_d.add_paragraph()
    p_body.text = "The initial brief prose stated '21,500 MAD (19,500 + 2,000)' due to a clerical typo. Our engine recalculation mathematically proves that (9.95 - 9.91) * 500,000 = 20,000 MAD (not 19,500 MAD), establishing the true figure of 22,000 MAD. Demonstrates true engineering judgment beyond blind transcription."
    p_body.font.size = Pt(10)
    p_body.font.color.rgb = C_NAVY
    p_body.space_before = Pt(4)

    set_notes(s5, "Comparing FIFO and CMP on the reference example: FIFO yields +22,000 MAD while CMP yields +18,000 MAD. Importantly, we detected a clerical typo in the source specification which wrote 21,500 MAD (19,500 + 2,000). Our recalculated figure proves 20,000 + 2,000 = 22,000 MAD, documented in Section 8.2.")

    # =========================================================================
    # SLIDE 6: 05 / RECONCILIATION — THE 4 REAL BANKING BREAKS (OS6)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "05", "Front-to-Back Reconciliation — 4 Banking Breaks")

    # 4 Break Cards (2x2 grid)
    b_w = Inches(5.75)
    b_h = Inches(2.05)

    breaks_data = [
        ("⚠️  1. Rate Slippage Break (BMCE Bank of Africa)", "FO booked EUR/MAD @ 10.7950, BO confirmed @ 10.7985 (+35 pips).\nFinancial Impact: -2,450 MAD fee discrepancy between trader ticket and bank SWIFT.", C_AMBER),
        ("❌  2. Nominal Mismatch Break (Attijariwafa Bank)", "FO booked 1,500,000 USD, BO settled only 1,400,000 USD (-100k USD cut).\nRoot Cause: Counterparty intraday credit ceiling was capped during business hours.", C_ROSE),
        ("⏳  3. Unconfirmed SWIFT MT300 (BNP Paribas Paris)", "Deal entered at 17:15 near close; no confirmation received before cutoff.\nOperational Risk: Unhedged overnight currency exposure with zero legal affirmation.", C_ROSE),
        ("👻  4. Orphan Settlement Record (Barclays Bank London)", "Back-Office received a confirmed MT300 settlement for 300k GBP with NO FO ticket.\nCompliance Flag: Detects unrecorded dealer trades or missing Front-Office booking.", C_BLUE)
    ]
    for i, (h_s, b_s, col) in enumerate(breaks_data):
        cx = Inches(0.8 + (i % 2) * 5.95)
        cy = Inches(1.5 + (i // 2) * 2.25)
        add_card(s6, cx, cy, b_w, b_h)
        tb = s6.shapes.add_textbox(cx + Inches(0.25), cy + Inches(0.18), b_w - Inches(0.5), b_h - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = h_s
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = col

        p_b = tf.add_paragraph()
        p_b.text = b_s
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = C_NAVY
        p_b.space_before = Pt(4)

    # Bottom Summary Bar
    sum_card = add_card(s6, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.85), bg=RGBColor(241, 245, 249))
    tb_s = s6.shapes.add_textbox(Inches(1.1), Inches(6.15), Inches(11.1), Inches(0.65))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "Audit Summary: 13 Clean Matched Trades (OCP, Renault, BCP) · 4 Isolated Discrepancies · 1-Click Excel Audit Package (.xlsx)"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NAVY

    set_notes(s6, "Our Front-to-Back Reconciliation engine (OS6) compares Front-Office tickets with Back-Office SWIFT MT300 records. We simulate 4 authentic banking breaks: rate slippage, nominal break from credit cap, unconfirmed SWIFT trade, and orphan settlement, exporting full audit reports to Excel.")

    # =========================================================================
    # SLIDE 7: 06 / PIPELINE ORCHESTRATION — APACHE AIRFLOW DAG
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "06", "Pipeline Orchestration — Apache Airflow DAG")

    # Top CRON Banner
    top_c = add_card(s7, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.9), bg=RGBColor(241, 245, 249))
    tb_top = s7.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(0.7))
    tf_top = tb_top.text_frame
    tf_top.word_wrap = True
    p = tf_top.paragraphs[0]
    p.text = "Scheduled at 15:35 Daily (CRON: 35 15 * * 1-5) — Triggered 5 Minutes After Bank Al-Maghrib Fixing"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = C_BLUE

    # 5 Sequential DAG Tasks (DermaGuard Horizontal Step style)
    dag_w = Inches(2.2)
    dag_h = Inches(3.2)
    dag_tasks = [
        ("Task 1", "Sensor", "wait_for_bam_fixing", "Polls BAM fixing table.\nRetries with exponential backoff if publication delayed.", C_BLUE),
        ("Task 2", "Ingestion", "ingest_eod_trades", "PythonOperator ingests FO tickets and BO SWIFT records into Bronze.", C_TEAL),
        ("Task 3", "dbt Mart", "dbt_run_models", "BashOperator executes fact_trade_mtm SQL transformations.", C_AMBER),
        ("Task 4", "Quality Gate", "dbt_test_integrity", "Circuit-breaker test: Halts DAG if BR1–BR4 or uniqueness fails.", C_ROSE),
        ("Task 5", "Calculation", "compute_pnl_engine", "Executes Python FIFO queue, updates daily_pnl & generates Excel.", C_EMERALD)
    ]
    for i, (t_num, t_type, t_name, t_desc, col) in enumerate(dag_tasks):
        dx = Inches(0.8 + i * 2.38)
        add_card(s7, dx, Inches(2.6), dag_w, dag_h)
        tb = s7.shapes.add_textbox(dx + Inches(0.15), Inches(2.75), dag_w - Inches(0.3), dag_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = t_num
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = col

        p_typ = tf.add_paragraph()
        p_typ.text = t_type
        p_typ.font.size = Pt(9.5)
        p_typ.font.bold = True
        p_typ.font.color.rgb = C_MUTED
        p_typ.space_before = Pt(2)

        p_nam = tf.add_paragraph()
        p_nam.text = t_name
        p_nam.font.name = FONT_MONO
        p_nam.font.size = Pt(9)
        p_nam.font.bold = True
        p_nam.font.color.rgb = C_NAVY
        p_nam.space_before = Pt(4)

        p_dsc = tf.add_paragraph()
        p_dsc.text = t_desc
        p_dsc.font.size = Pt(8.5)
        p_dsc.font.color.rgb = C_NAVY
        p_dsc.space_before = Pt(4)

    # Bottom Guarantees Card
    g_card = add_card(s7, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), bg=CARD_BG)
    tb_g = s7.shapes.add_textbox(Inches(1.1), Inches(6.12), Inches(11.1), Inches(0.65))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True
    p = tf_g.paragraphs[0]
    p.text = "Production Guarantees: Idempotent re-runs without duplicate rows · Circuit-breaker error handling · Automatic SQLite fallback."
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NAVY

    set_notes(s7, "Our workflow is fully automated with Apache Airflow. Triggered daily at 15:35, the DAG begins with a BAM fixing sensor, ingests trades, runs dbt SQL models, executes integrity tests as circuit-breakers, and finishes with Python FIFO P&L calculation and Excel generation.")

    # =========================================================================
    # SLIDE 8: 07 / AUTOMATED TESTING & DATA CONTRACTS (10/10 PASS)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "07", "Automated Testing & Data Contracts (10/10 PASS)")

    c_w = Inches(5.75)
    c_h = Inches(5.4)

    # Left: Pytest Suite
    c_l = add_card(s8, Inches(0.8), Inches(1.5), c_w, c_h)
    tb_l = s8.shapes.add_textbox(Inches(1.05), Inches(1.7), c_w - Inches(0.5), c_h - Inches(0.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Pytest Automated Suite (10 Tests in 0.51s)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_EMERALD

    p_tests = [
        ("test_individual_trade_mtm", "MTM accuracy on USD/MAD, EUR/MAD, and cross EUR/USD."),
        ("test_realized_and_unrealized_split", "Exact split: +20k Realized / +59,460 Latent."),
        ("test_supporting_example_fifo_vs_cmp", "Section 4 validation: +22,000 FIFO vs +18,000 CMP."),
        ("test_business_rule_br3_identity", "Total = Realized + Latent validated across 100 random deals."),
        ("test_reconciliation_exact_matches", "0 false positives on clean matched trades."),
        ("test_reconciliation_rate_slippage", "Pip-level variance detection and MAD calculation."),
        ("test_reconciliation_nominal_mismatch", "Partial fill and credit ceiling break detection."),
        ("test_reconciliation_unconfirmed_trade", "Unconfirmed SWIFT MT300 ghost trade flagging.")
    ]
    for name, desc in p_tests:
        p = tf_l.add_paragraph()
        p.text = f"✔ {name}"
        p.font.name = FONT_MONO
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD
        p.space_before = Pt(3)
        p_d = tf_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = C_NAVY

    # Right: dbt Rules
    c_r = add_card(s8, Inches(6.8), Inches(1.5), c_w, c_h)
    tb_r = s8.shapes.add_textbox(Inches(7.05), Inches(1.7), c_w - Inches(0.5), c_h - Inches(0.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "dbt Automated Data Contracts (Report Table 8.1)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_BLUE

    dbt_rules = [
        ("BR1: BAM Rate Conversion", "Guarantees all foreign pairs convert to MAD; rejects null or negative rates.", "PASS"),
        ("BR2: Rate Revaluation", "Verifies all open tickets are revalued at today's official BAM fixing.", "PASS"),
        ("BR3: Accounting Equality", "Singular SQL test: SELECT * WHERE ROUND(total) != ROUND(realized + latent). Returns 0 rows.", "PASS"),
        ("BR4: Chronological Lot Order", "Verifies queue unwinding cannot consume a newer lot before earlier timestamped lots.", "PASS"),
        ("source_ref Uniqueness", "Constraint preventing duplicate trade bookings in database.", "PASS")
    ]
    for r_name, r_desc, r_stat in dbt_rules:
        p = tf_r.add_paragraph()
        p.text = f"• {r_name}  [{r_stat}]"
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_NAVY
        p.space_before = Pt(6)
        p_d = tf_r.add_paragraph()
        p_d.text = r_desc
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = C_MUTED

    set_notes(s8, "Our software quality is proven by 10 automated unit tests passing in 0.51s, paired with dbt singular tests that enforce our business rules (BR1 to BR4). The physical CHECK constraint in PostgreSQL guarantees that Total = Realized + Unrealized at all times.")

    # =========================================================================
    # SLIDE 9: 08 / OPERATIONAL WEB TERMINAL — STREAMLIT TRADING DESK
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "08", "Operational Web Terminal — Streamlit Trading Desk")

    # 4 Operational Capability Cards (DermaGuard style)
    c_w = Inches(5.75)
    c_h = Inches(2.5)
    caps = [
        ("🟢  Live BAM Fixing Strip & Basket", "Real-time USD/MAD, EUR/MAD, and central basket (60% EUR + 40% USD) tracking with +/- 5% peg band boundaries.", C_EMERALD),
        ("🔍  Trade Blotter with Transparent Math", "Complete breakdown of each trade showing the exact arithmetic formula in MAD (no black-box calculations).", C_BLUE),
        ("⚡  Pre-Trade Deal Ticket Simulator", "Simulate the instant MTM, MAD exposure, and balance sheet impact of an order before committing it to the book.", C_AMBER),
        ("🛡️  FO / BO Dispute Hub & Excel Export", "One-click filtering of breaks (slippage, nominal, unconfirmed), discrepancy tags, and instant .xlsx audit download.", C_PURPLE)
    ]
    for i, (h_s, b_s, col) in enumerate(caps):
        cx = Inches(0.8 + (i % 2) * 5.95)
        cy = Inches(1.5 + (i // 2) * 2.7)
        add_card(s9, cx, cy, c_w, c_h)
        tb = s9.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.25), c_w - Inches(0.6), c_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = h_s
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = col

        p_b = tf.add_paragraph()
        p_b.text = b_s
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = C_NAVY
        p_b.space_before = Pt(8)

    # Bottom Terminal Status Card
    bot_t = add_card(s9, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), bg=RGBColor(241, 245, 249))
    tb_bt = s9.shapes.add_textbox(Inches(1.1), Inches(6.15), Inches(11.1), Inches(0.6))
    tf_bt = tb_bt.text_frame
    tf_bt.word_wrap = True
    p = tf_bt.paragraphs[0]
    p.text = "Operational Platform: Streamlit Web Terminal · Plotly Visualization · Dual PostgreSQL/SQLite Persistence · Zero Runtime Crashes."
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NAVY

    set_notes(s9, "Our Streamlit dashboard serves as a functional trading terminal: live BAM fixings strip, transparent mathematical formulas on every trade, pre-trade deal simulator, and Front-to-Back dispute hub with 1-click Excel export.")

    # =========================================================================
    # SLIDE 10: 09 / CONCLUSION & EVALUATION ALIGNMENT (DERMAGUARD SUMMARY STYLE)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "09", "Conclusion & Evaluation Alignment")

    # Left: Evaluation Criteria Table from Report Table 10.3
    c_l = add_card(s10, Inches(0.8), Inches(1.5), Inches(6.8), Inches(4.5))
    tb_l = s10.shapes.add_textbox(Inches(1.05), Inches(1.65), Inches(6.3), Inches(4.2))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Alignment with Evaluation Criteria (Report Table 10.3)"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_BLUE

    crit_rows = [
        ("25% Business-Domain Mastery:", "BAM peg compliance, MTM, Realized/Latent, and brief typo audit correction."),
        ("20% Design Quality:", "Medallion layered pipeline (Bronze/Silver/Gold), ERD with logical lookup."),
        ("25% Development Quality:", "10/10 automated tests, dbt data contracts, dual PostgreSQL/SQLite storage."),
        ("15% Relevance of Results:", "Validated benchmark (+79,460 MAD) and realistic FO/BO reconciliation breaks."),
        ("10% Documentation & Report:", "Comprehensive 33-page thesis report and reproducible Docker environment."),
        ("5% Oral Defense:", "This presentation!")
    ]
    for h_s, b_s in crit_rows:
        p = tf_l.add_paragraph()
        p.text = f"✔ {h_s} {b_s}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_NAVY
        p.space_before = Pt(4)

    # Right: Summary & Key Takeaway
    c_r = add_card(s10, Inches(7.8), Inches(1.5), Inches(4.733), Inches(4.5))
    tb_r = s10.shapes.add_textbox(Inches(8.05), Inches(1.65), Inches(4.233), Inches(4.2))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Project Summary & Perspectives"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = C_EMERALD

    sum_items = [
        ("Accomplished:", "Delivered a fully compliant, tested, and automated FX P&L platform replacing manual Excel sheets at DXC."),
        ("Future Perspectives:", "Real-time tick streaming with Apache Kafka and direct SWIFT MT300 message network ingestion.")
    ]
    for h_s, b_s in sum_items:
        p = tf_r.add_paragraph()
        p.text = f"• {h_s} {b_s}"
        p.font.size = Pt(10)
        p.font.color.rgb = C_NAVY
        p.space_before = Pt(8)

    p_thx = tf_r.add_paragraph()
    p_thx.text = "\nThank you for your attention.\nQuestions & Discussion."
    p_thx.font.bold = True
    p_thx.font.size = Pt(14)
    p_thx.font.color.rgb = C_NAVY
    p_thx.space_before = Pt(14)

    # Bottom Quote Card (Exact DermaGuard closing quote style)
    q_card = add_card(s10, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.85), bg=RGBColor(241, 245, 249))
    tb_q = s10.shapes.add_textbox(Inches(1.1), Inches(6.25), Inches(11.1), Inches(0.65))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    p = tf_q.paragraphs[0]
    p.text = '"The goal is not to build the most complex system. It is to build a platform where every calculation is mathematically audited, every break is explained, and the bank is 100% compliant with Bank Al-Maghrib."'
    p.font.name = FONT_MAIN
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    set_notes(s10, "To conclude: this project bridges financial engineering with modern data engineering. All evaluation criteria have been met and validated. I thank Prof. Zinedine and Mr. Gourroum for their guidance, and I welcome your questions.")

    # Save presentation
    out_path = PRES_DIR / "FX_Trading_Room_PFE_Defense_DermaGuard_Style.pptx"
    prs.save(out_path)
    print(f"DermaGuard-style presentation created successfully at: {out_path}")


if __name__ == "__main__":
    build_deck()
